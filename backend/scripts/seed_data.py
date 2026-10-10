"""Convert, validate, and optionally seed Mission 1 records to Firestore."""

from __future__ import annotations

import argparse
import math
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from statistics import fmean
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = (
    PROJECT_ROOT.parent / "branding-trend-analysis" / "data" / "datalab.xlsx"
)
SHEET_NAME = "개요"
SOURCE_RANGE = "A1:F193"
HEADER_ROW = 7
FIRST_DATA_ROW = 8
LAST_DATA_ROW = 193
EXPECTED_WEEK_COUNT = 186
EXPECTED_RECORD_COUNT = 558
MEMO = "네이버 데이터랩 주간 검색 관심도"
KEYWORDS = ("퍼스널브랜딩", "이미지메이킹", "스피치")
KEYWORD_SLUGS = {
    "퍼스널브랜딩": "personal-branding",
    "이미지메이킹": "image-making",
    "스피치": "speech",
}
FIRESTORE_COLLECTION = "data"
FIRESTORE_BATCH_SIZE = 400
FIRESTORE_BATCH_MAX_WRITES = 500
EXPECTED_HEADERS = (
    "날짜",
    "퍼스널브랜딩",
    "날짜",
    "이미지메이킹",
    "날짜",
    "스피치",
)
COLUMN_GROUPS = (
    (1, 2, "퍼스널브랜딩"),
    (3, 4, "이미지메이킹"),
    (5, 6, "스피치"),
)
REFERENCE_STATISTICS = {
    "퍼스널브랜딩": {
        "average": 20.930603,
        "minimum": 4.69745,
        "maximum": 46.75557,
    },
    "이미지메이킹": {
        "average": 9.533696,
        "minimum": 3.18471,
        "maximum": 19.62579,
    },
    "스피치": {
        "average": 65.386569,
        "minimum": 37.04219,
        "maximum": 100.0,
    },
}
STATISTICS_ABS_TOLERANCE = 1e-6


class SeedDataValidationError(RuntimeError):
    """Raised when the source workbook cannot produce trusted seed records."""


def _parse_date(value: Any, row: int, column: int) -> date:
    if value is None or value == "":
        raise SeedDataValidationError(
            f"원본 {row}행 {column}열의 날짜가 비어 있습니다."
        )

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass

    raise SeedDataValidationError(
        f"원본 {row}행 {column}열을 YYYY-MM-DD 날짜로 변환할 수 없습니다."
    )


def _parse_float(value: Any, row: int, column: int) -> float:
    if value is None or value == "":
        raise SeedDataValidationError(
            f"원본 {row}행 {column}열의 관심도 값이 비어 있습니다."
        )
    if isinstance(value, bool):
        raise SeedDataValidationError(
            f"원본 {row}행 {column}열의 관심도 값은 숫자여야 합니다."
        )

    try:
        converted = float(value)
    except (TypeError, ValueError):
        raise SeedDataValidationError(
            f"원본 {row}행 {column}열의 관심도 값을 float로 변환할 수 없습니다."
        ) from None

    if not math.isfinite(converted):
        raise SeedDataValidationError(
            f"원본 {row}행 {column}열의 관심도 값이 유한한 숫자가 아닙니다."
        )
    return converted


def _validate_records(records: list[dict[str, str | float]]) -> dict[str, Any]:
    required_fields = {"date", "keyword", "value", "memo"}
    for index, record in enumerate(records, start=1):
        missing_fields = required_fields - record.keys()
        if missing_fields:
            missing = ", ".join(sorted(missing_fields))
            raise SeedDataValidationError(
                f"레코드 {index}에 필수 필드가 없습니다: {missing}"
            )

        date_value = record["date"]
        keyword = record["keyword"]
        value = record["value"]
        memo = record["memo"]

        if not isinstance(date_value, str):
            raise SeedDataValidationError(f"레코드 {index}의 date가 문자열이 아닙니다.")
        try:
            parsed_date = date.fromisoformat(date_value)
        except ValueError as exc:
            raise SeedDataValidationError(
                f"레코드 {index}의 date 형식이 YYYY-MM-DD가 아닙니다."
            ) from exc
        if parsed_date.isoformat() != date_value:
            raise SeedDataValidationError(
                f"레코드 {index}의 date 형식이 YYYY-MM-DD가 아닙니다."
            )
        if keyword not in KEYWORDS:
            raise SeedDataValidationError(
                f"레코드 {index}에 허용되지 않은 keyword가 있습니다."
            )
        if type(value) is not float or not math.isfinite(value):
            raise SeedDataValidationError(
                f"레코드 {index}의 value는 유한한 float여야 합니다."
            )
        if memo is not None and not isinstance(memo, str):
            raise SeedDataValidationError(
                f"레코드 {index}의 memo는 문자열 또는 null이어야 합니다."
            )

    if len(records) != EXPECTED_RECORD_COUNT:
        raise SeedDataValidationError(
            f"총 레코드는 {EXPECTED_RECORD_COUNT}개여야 하지만 {len(records)}개입니다."
        )

    combinations = [(record["date"], record["keyword"]) for record in records]
    duplicate_count = len(combinations) - len(set(combinations))
    if duplicate_count:
        raise SeedDataValidationError(
            f"중복된 날짜+키워드 조합이 {duplicate_count}개 있습니다."
        )

    keyword_counts = Counter(str(record["keyword"]) for record in records)
    expected_counts = {keyword: EXPECTED_WEEK_COUNT for keyword in KEYWORDS}
    if dict(keyword_counts) != expected_counts:
        raise SeedDataValidationError(
            f"키워드별 레코드 수가 예상과 다릅니다: {dict(keyword_counts)}"
        )

    unique_date_strings = list(dict.fromkeys(str(record["date"]) for record in records))
    unique_dates = [date.fromisoformat(value) for value in unique_date_strings]
    if len(unique_dates) != EXPECTED_WEEK_COUNT:
        raise SeedDataValidationError(
            f"주간 날짜는 {EXPECTED_WEEK_COUNT}개여야 하지만 {len(unique_dates)}개입니다."
        )
    if unique_dates != sorted(unique_dates):
        raise SeedDataValidationError("주간 날짜가 오름차순이 아닙니다.")

    intervals = [
        (current - previous).days
        for previous, current in zip(unique_dates, unique_dates[1:])
    ]
    if any(interval != 7 for interval in intervals):
        raise SeedDataValidationError("주간 날짜 간격에 7일이 아닌 구간이 있습니다.")

    statistics = {}
    reference_matches = {}
    for keyword in KEYWORDS:
        values = [
            float(record["value"])
            for record in records
            if record["keyword"] == keyword
        ]
        observed = {
            "average": fmean(values),
            "minimum": min(values),
            "maximum": max(values),
        }
        statistics[keyword] = observed
        reference_matches[keyword] = {
            name: math.isclose(
                observed[name],
                REFERENCE_STATISTICS[keyword][name],
                rel_tol=0.0,
                abs_tol=STATISTICS_ABS_TOLERANCE,
            )
            for name in ("average", "minimum", "maximum")
        }

    if not all(
        matched
        for keyword_result in reference_matches.values()
        for matched in keyword_result.values()
    ):
        raise SeedDataValidationError("계산 통계가 검증된 기준값과 일치하지 않습니다.")

    return {
        "sheet_name": SHEET_NAME,
        "source_range": SOURCE_RANGE,
        "header_range": "A7:F7",
        "data_range": "A8:F193",
        "week_count": len(unique_dates),
        "record_count": len(records),
        "date_columns_identical": True,
        "dates_ascending": True,
        "all_intervals_seven_days": True,
        "duplicate_date_keyword_count": duplicate_count,
        "missing_value_count": 0,
        "conversion_failure_count": 0,
        "all_values_float": all(type(record["value"]) is float for record in records),
        "all_values_finite": all(
            math.isfinite(float(record["value"])) for record in records
        ),
        "keyword_counts": dict(keyword_counts),
        "start_date": unique_dates[0].isoformat(),
        "end_date": unique_dates[-1].isoformat(),
        "statistics": statistics,
        "reference_matches": reference_matches,
        "statistics_abs_tolerance": STATISTICS_ABS_TOLERANCE,
    }


def make_document_id(date_value: str, keyword: str) -> str:
    """Return the same Firestore ID for the same date and keyword."""
    try:
        parsed_date = date.fromisoformat(date_value)
    except (TypeError, ValueError) as exc:
        raise SeedDataValidationError("문서 ID의 날짜 형식이 올바르지 않습니다.") from exc
    if parsed_date.isoformat() != date_value:
        raise SeedDataValidationError("문서 ID의 날짜 형식이 YYYY-MM-DD가 아닙니다.")

    try:
        keyword_slug = KEYWORD_SLUGS[keyword]
    except KeyError as exc:
        raise SeedDataValidationError("문서 ID에 허용되지 않은 키워드가 있습니다.") from exc
    return f"{date_value}_{keyword_slug}"


def write_seed_records(
    records: list[dict[str, str | float]],
    client: Any,
    batch_size: int = FIRESTORE_BATCH_SIZE,
) -> dict[str, Any]:
    """Validate records, then write them using deterministic IDs and batches.

    Re-running this function overwrites the same document IDs with ``set``;
    it does not create additional documents for the same date and keyword.
    """
    validation = _validate_records(records)
    if not 1 <= batch_size <= FIRESTORE_BATCH_MAX_WRITES:
        raise SeedDataValidationError(
            f"batch_size는 1~{FIRESTORE_BATCH_MAX_WRITES} 범위여야 합니다."
        )

    from firebase_admin import firestore

    collection = client.collection(FIRESTORE_COLLECTION)
    batch_sizes: list[int] = []
    for start in range(0, len(records), batch_size):
        chunk = records[start : start + batch_size]
        batch = client.batch()
        for record in chunk:
            document_id = make_document_id(
                str(record["date"]), str(record["keyword"])
            )
            payload = {
                "date": record["date"],
                "keyword": record["keyword"],
                "value": record["value"],
                "memo": record["memo"],
                "created_at": firestore.SERVER_TIMESTAMP,
                "updated_at": firestore.SERVER_TIMESTAMP,
            }
            batch.set(collection.document(document_id), payload)
        batch.commit()
        batch_sizes.append(len(chunk))

    return {
        "write_count": len(records),
        "batch_count": len(batch_sizes),
        "batch_sizes": batch_sizes,
        "validation": validation,
    }


def verify_firestore_seed(client: Any) -> dict[str, Any]:
    """Read the data collection and verify the complete seeded dataset."""
    snapshots = list(client.collection(FIRESTORE_COLLECTION).stream())
    records: list[dict[str, str | float]] = []
    document_id_mismatches = 0

    for snapshot in snapshots:
        stored = snapshot.to_dict()
        record = {
            "date": stored.get("date"),
            "keyword": stored.get("keyword"),
            "value": stored.get("value"),
            "memo": stored.get("memo"),
        }
        records.append(record)
        try:
            expected_id = make_document_id(record["date"], record["keyword"])
        except SeedDataValidationError:
            document_id_mismatches += 1
        else:
            if snapshot.id != expected_id:
                document_id_mismatches += 1

    keyword_order = {keyword: index for index, keyword in enumerate(KEYWORDS)}
    records.sort(
        key=lambda item: (
            str(item["date"]),
            keyword_order.get(str(item["keyword"]), len(KEYWORDS)),
        )
    )
    validation = _validate_records(records)
    if document_id_mismatches:
        raise SeedDataValidationError(
            f"deterministic 문서 ID 불일치: {document_id_mismatches}개"
        )

    return {
        **validation,
        "firestore_document_count": len(snapshots),
        "document_id_mismatches": document_id_mismatches,
    }


def run_seed(
    records: list[dict[str, str | float]],
    *,
    write: bool = False,
    client: Any | None = None,
) -> dict[str, Any]:
    """Validate only by default; write and verify only when explicitly enabled."""
    validation = _validate_records(records)
    if not write:
        return {
            "mode": "dry-run",
            "expected_document_count": EXPECTED_RECORD_COUNT,
            "write_count": 0,
            "validation": validation,
        }

    if client is None:
        from backend.app.database import get_firestore_client

        client = get_firestore_client()

    write_result = write_seed_records(records, client)
    post_write_validation = verify_firestore_seed(client)
    return {
        "mode": "write",
        "expected_document_count": EXPECTED_RECORD_COUNT,
        "write_count": write_result["write_count"],
        "batch_count": write_result["batch_count"],
        "batch_sizes": write_result["batch_sizes"],
        "validation": validation,
        "post_write_validation": post_write_validation,
    }


def convert_workbook_to_records(
    workbook_path: str | Path = DEFAULT_WORKBOOK_PATH,
) -> tuple[list[dict[str, str | float]], dict[str, Any]]:
    """Read the workbook in normal mode, then convert and validate 558 records."""

    source_path = Path(workbook_path)
    if not source_path.is_file():
        raise SeedDataValidationError(f"원본 엑셀 파일을 찾을 수 없습니다: {source_path}")

    workbook = load_workbook(source_path, read_only=False, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SeedDataValidationError(f"원본에 '{SHEET_NAME}' 시트가 없습니다.")
        worksheet = workbook[SHEET_NAME]

        headers = tuple(
            worksheet.cell(row=HEADER_ROW, column=column).value
            for column in range(1, 7)
        )
        if headers != EXPECTED_HEADERS:
            raise SeedDataValidationError(
                f"7행 헤더가 예상 구조와 다릅니다: {headers}"
            )

        records: list[dict[str, str | float]] = []
        for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            parsed_dates = [
                _parse_date(worksheet.cell(row=row, column=date_column).value, row, date_column)
                for date_column, _, _ in COLUMN_GROUPS
            ]
            if len(set(parsed_dates)) != 1:
                raise SeedDataValidationError(
                    f"원본 {row}행의 세 날짜 열이 서로 다릅니다."
                )

            date_text = parsed_dates[0].isoformat()
            for _, value_column, keyword in COLUMN_GROUPS:
                records.append(
                    {
                        "date": date_text,
                        "keyword": keyword,
                        "value": _parse_float(
                            worksheet.cell(row=row, column=value_column).value,
                            row,
                            value_column,
                        ),
                        "memo": MEMO,
                    }
                )
    finally:
        workbook.close()

    validation = _validate_records(records)
    return records, validation


def _build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="미션 1 Excel 데이터를 Firestore seed 레코드로 변환합니다."
    )
    parser.add_argument(
        "--write",
        action="store_true",
        help="검증된 558개 레코드를 Firestore에 실제 적재합니다.",
    )
    return parser


def main() -> None:
    args = _build_argument_parser().parse_args()
    records, validation = convert_workbook_to_records()
    result = run_seed(records, write=args.write)

    print(f"Firestore seed 실행 모드: {result['mode']}")
    print(f"- 시트/범위: {validation['sheet_name']} {validation['source_range']}")
    print(f"- 주간 날짜: {validation['week_count']}개")
    print(f"- 총 레코드: {len(records)}개")
    print(f"- 기간: {validation['start_date']} ~ {validation['end_date']}")
    print(f"- 키워드별 개수: {validation['keyword_counts']}")
    print("- 기준 통계 대조: 모두 일치")
    print(f"- 예상 Firestore 문서 수: {result['expected_document_count']}개")

    if args.write:
        post_validation = result["post_write_validation"]
        print(f"- 실제 Firestore write: {result['write_count']}건")
        print(f"- batch 분할: {result['batch_sizes']}")
        print(
            "- Firestore 사후 검증 문서 수: "
            f"{post_validation['firestore_document_count']}개"
        )
        print("- Firestore 사후 검증: 통과")
    else:
        print("- 실제 Firestore write: 0건")


if __name__ == "__main__":
    main()
